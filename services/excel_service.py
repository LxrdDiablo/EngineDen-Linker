from openpyxl import load_workbook

from models.product import Product


class ExcelService:
    """
    Reads and writes EngineDen workbooks.
    """

    def load_products(self, filename):

        workbook = load_workbook(filename)

        sheet = workbook.active

        products = []

        for row in range(2, sheet.max_row + 1):

            description = sheet.cell(row=row, column=1).value

            if not description:
                continue

            products.append(

                Product(
                    row_number=row,
                    description=str(description).strip()
                )

            )

        workbook.close()

        return products

    def save_products(self, input_file, output_file, products):

        workbook = load_workbook(input_file)

        sheet = workbook.active

        headers = {

            "Matched Product": 2,
            "Confidence": 3,
            "EngineDen URL": 4,
            "Image URL": 5,
            "Status": 6

        }

        for title, column in headers.items():

            sheet.cell(row=1, column=column).value = title

        for product in products:

            row = product.row_number

            sheet.cell(row=row, column=2).value = product.matched_name

            sheet.cell(row=row, column=3).value = product.confidence

            sheet.cell(row=row, column=4).value = product.engineden_url

            sheet.cell(row=row, column=5).value = product.image_url

            sheet.cell(row=row, column=6).value = product.status

        workbook.save(output_file)

        workbook.close()