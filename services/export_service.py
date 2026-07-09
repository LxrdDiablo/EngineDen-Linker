import os
from datetime import datetime

from openpyxl import load_workbook


class ExportService:
    """
    Exports EngineDen results without overwriting
    the original workbook columns.
    """

    def export(self, input_file, output_folder, products):

        workbook = load_workbook(input_file)

        sheet = workbook.active

        start_column = sheet.max_column + 1

        headers = [

            "Matched Product",

            "Confidence",

            "EngineDen URL",

            "Image URL",

            "Status"

        ]

        for offset, title in enumerate(headers):

            sheet.cell(
                row=1,
                column=start_column + offset
            ).value = title

        for product in products:

            row = product.row_number

            values = [

                product.matched_name,

                product.confidence,

                product.engineden_url,

                product.image_url,

                product.status

            ]

            for offset, value in enumerate(values):

                sheet.cell(
                    row=row,
                    column=start_column + offset
                ).value = value

        filename = (
            "EngineDen_Output_"
            + datetime.now().strftime("%Y%m%d_%H%M%S")
            + ".xlsx"
        )

        output_file = os.path.join(
            output_folder,
            filename
        )

        workbook.save(output_file)

        workbook.close()

        return output_file